"""Polish the existing report without requiring the original template again."""
import ast
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_SECTION_START
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree as ET

root=Path(__file__).resolve().parents[1]
path=root/'docs/CEP_Report.docx'
doc=Document(path)
anchors={15:'Seasonal naive forecasting repeats',22:'Training covers 2012-2022',31:'All 10 automated tests passed',24:'Algorithm steps are shown below'}
body=next(p for p in doc.paragraphs if p.style.name=='Body Text')
existing='\n'.join(p.text for p in doc.paragraphs)
anchor=None
for node in ast.parse((root/'tools/build_report.py').read_text()).body:
    if not isinstance(node,ast.Assign) or not any(isinstance(t,ast.Name) and t.id=='theory' for t in node.targets):continue
    call=node.value
    if not isinstance(call,ast.Call) or not isinstance(call.func,ast.Name) or call.func.id!='insert_after':continue
    if isinstance(call.args[0],ast.Subscript):
        index=ast.literal_eval(call.args[0].slice)
        anchor=next(p for p in doc.paragraphs if p.text.startswith(anchors[index]))
    text=ast.literal_eval(call.args[1])
    if text in existing:continue
    el=deepcopy(body._p)
    for child in list(el):
        if child.tag!=qn('w:pPr'):el.remove(child)
    anchor._p.addnext(el)
    anchor=Paragraph(el,anchor._parent)
    bold=any(k.arg=='bold' and ast.literal_eval(k.value) for k in call.keywords)
    anchor.add_run(text).bold=bold
    anchor.paragraph_format.keep_with_next=bold

for p in list(doc.paragraphs):
    p.paragraph_format.page_break_before=False
    p.paragraph_format.space_before=Pt(2)
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.widow_control=True
    if p.style.name.startswith('Heading') or (p.text and len(p.text)<120 and all(r.bold for r in p.runs if r.text)):
        p.paragraph_format.keep_with_next=True
    if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:pict|.//w:sectPr'):
        p._p.getparent().remove(p._p)
for section in list(doc.sections)[1:]:section.start_type=WD_SECTION_START.CONTINUOUS
for section in doc.sections:
    section.bottom_margin=Pt(65)
for table in doc.tables:
    for i,row in enumerate(table.rows):
        props=row._tr.get_or_add_trPr()
        for old in list(props.findall(qn('w:tblHeader'))):props.remove(old)
        if i==0:props.append(OxmlElement('w:tblHeader'))
        if len(table.rows)<=8:
            for cell in row.cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=i<len(table.rows)-1
doc.save(path)
with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
for name in list(parts):
    if name.startswith('word/header') and name.endswith('.xml'):
        xml=ET.fromstring(parts[name])
        for drawing in list(xml.iter(qn('w:drawing'))):drawing.getparent().remove(drawing)
        parts[name]=ET.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(path,'w',ZIP_DEFLATED) as z:
    for name,content in parts.items():z.writestr(name,content)
print('Report theory and continuous layout updated.')
