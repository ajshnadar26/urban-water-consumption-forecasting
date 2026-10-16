"""Build from the retained DAV template, preserving its styles and page furniture.
Usage: python tools/build_report.py --template path/to/template.docx
"""
import argparse, copy, hashlib, json, sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from pygments import lex
from pygments.lexers import PythonLexer
from pygments.token import Token
from docx.shared import RGBColor

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
parser=argparse.ArgumentParser();parser.add_argument('--template',required=True);args=parser.parse_args()
template=Path(args.template)
d=Document(template); original=list(d.paragraphs); source_tables=list(d.tables)
summary=json.loads((ROOT/'results/summary.json').read_text())
metrics=pd.read_csv(ROOT/'results/metrics.csv')
forecast=pd.read_csv(ROOT/'results/future_forecast.csv')
anomalies=pd.read_csv(ROOT/'results/anomalies.csv')

def replace(index,text):
    p=original[index]
    # Preserve paragraph and first-run formatting, replace only the editable text.
    runs=p.runs
    if runs:
        runs[0].text=text
        for run in runs[1:]:run.text=''
    else:p.add_run(text)
    return p

def body(index,text):
    p=replace(index,text);p.style='Body Text'
    num=p._p.pPr.find(qn('w:numPr')) if p._p.pPr is not None else None
    if num is not None:p._p.pPr.remove(num)
    p.paragraph_format.keep_with_next=False
    return p

def insert_after(anchor,text='',bold=False,pagebreak=False):
    el=copy.deepcopy(original[10]._p)
    for child in list(el):
        if child.tag!=qn('w:pPr'):el.remove(child)
    anchor._p.addnext(el);p=Paragraph(el,anchor._parent)
    p.paragraph_format.keep_with_next=False
    p.paragraph_format.space_after=Pt(6)
    run=p.add_run(text);run.bold=bold
    if pagebreak:p.paragraph_format.page_break_before=True
    return p

def page_before(index):original[index].paragraph_format.page_break_before=True

def code_after(anchor, source):
    for line in source.splitlines():
        p=insert_after(anchor);p.paragraph_format.space_after=Pt(2)
        for kind,value in lex(line,PythonLexer()):
            value=value.rstrip('\n')
            if not value:continue
            r=p.add_run(value);r.font.name='Consolas';r.font.size=Pt(9)
            color='1F2937'
            if kind in Token.Keyword:color='7C3AED'
            elif kind in Token.Literal.String:color='18743B'
            elif kind in Token.Literal.Number:color='B45309'
            elif kind in Token.Name.Builtin:color='005F99'
            elif kind in Token.Comment:color='64748B'
            r.font.color.rgb=RGBColor.from_string(color)
        anchor=p
    return anchor

def figure(anchor,path,caption,width=6.65):
    p=insert_after(anchor);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/path),width=Inches(width))
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=insert_after(p,caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p

def table_after(anchor,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers))
    anchor._p.addnext(t._tbl)
    t.autofit=False
    for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    style_table(t,widths)
    # Return a following body paragraph as the insertion anchor.
    el=copy.deepcopy(original[10]._p)
    for child in list(el):
        if child.tag!=qn('w:pPr'):el.remove(child)
    t._tbl.addnext(el)
    return Paragraph(el,anchor._parent)

def style_table(t,widths=None):
    if widths:
        for row in t.rows:
            for cell,width in zip(row.cells,widths):cell.width=Inches(width)
        for col,width in zip(t.columns,widths):col.width=Inches(width)
    props=t._tbl.tblPr
    borders=props.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');props.append(borders)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'000000');borders.append(item)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr()
            for existing in list(tcPr.findall(qn('w:shd'))):tcPr.remove(existing)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                val=OxmlElement('w:'+side);val.set(qn('w:w'),'70');val.set(qn('w:type'),'dxa');margins.append(val)
            tcPr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.keep_with_next=False
                for r in p.runs:r.font.size=Pt(9);r.bold=(i==0)

# Architecture is generated from the implemented modules.
fig,ax=plt.subplots(figsize=(10,3.5));ax.axis('off');ax.set_xlim(-.03,1.03);ax.set_ylim(0,1)
boxes=[(.01,.62,.27,.23,'Austin retail billing\nDaily ERA5 weather'),(.36,.62,.27,.23,'Validation and monthly ETL\nCalendar and sector mapping'),(.72,.62,.27,.23,'Train → validation → test\nARIMA and alternatives'),(.72,.13,.27,.23,'Forecasts and metrics\nCausal anomaly screening'),(.36,.13,.27,.23,'Sector shares and allocation\nFloors, priorities and caps'),(.01,.13,.27,.23,'Streamlit decision support\nCharts and CSV downloads')]
for i,(x,y,w,h,label) in enumerate(boxes):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='square,pad=.012',facecolor=['#E0F2FE','#DCFCE7','#F3E8FF','#FEF3C7','#FFE4E6','#CCFBF1'][i],edgecolor='#475569'))
    ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=10,color='black')
for a,b in [((.28,.73),(.36,.73)),((.63,.73),(.72,.73)),((.85,.62),(.85,.36)),((.72,.24),(.63,.24)),((.36,.24),(.28,.24))]:
    ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',color='black',lw=1.2))
fig.tight_layout();fig.savefig(ROOT/'results/figures/architecture.png',dpi=600,bbox_inches='tight');plt.close(fig)

replace(2,'DATA ANALYTICS AND VISUALIZATION DAV PROJECT REPORT')
replace(4,'Project Title: Urban Water Consumption Forecasting and Smart Resource Allocation\nGroup No.: __________    Faculty Name: __________________    Submission Date: __________\nStudent Names and Roll Numbers')
replace(5,'1. ____________________   2. ____________________   3. ____________________')
replace(6,'4. ____________________   5. ____________________   6. ____________________')
replace(7,'7. ____________________   8. ____________________   9. ____________________')
member_anchor=insert_after(original[7],'10. ____________________   11. ____________________   12. ____________________')
insert_after(member_anchor,'13. ____________________')
body(10,'This project develops a time-series decision-support system for urban water demand and constrained distribution. City of Austin residential and commercial billing snapshots are combined with ERA5 weather data from Open-Meteo. The analysis covers 156 monthly observations from January 2012 to December 2024. Data validation, seasonal exploration, rolling statistics and the Augmented Dickey-Fuller test support model development. ARIMA is compared with Holt-Winters, seasonal naive and weather/calendar regression using chronological validation and testing. Holt-Winters is selected on 2023 validation and achieves 250.18 MG MAE, 299.87 MG RMSE and 6.70 percent MAPE on the 2024 holdout. A causal residual screen identifies four unusual months for review. Priority and proportional allocation scenarios enforce supply and capacity limits and expose infeasible minimum service. An interactive dashboard connects forecasts, anomalies and recommended actions. Operational deployment requires verified supply, infrastructure and priority-facility data.')
body(12,'Municipal demand changes with season, weather, growth and consumer activity. Fixed distribution plans can produce shortages in some areas while leaving supply unused elsewhere. The engineering problem is to forecast demand, identify unusual consumption and recommend a transparent distribution strategy when aggregate supply or sector capacity is limited.')
body(13,'Objectives are to validate water and weather data, quantify seasonality and stationarity, compare forecasting models fairly, screen abnormal demand and evaluate alternative allocations. The scope is selected retail billing classes in Austin. Large-volume billing is an industrial proxy and municipal billing is a priority-area proxy. Supply limits, lifeline shares and network capacities are scenario assumptions. The 2025 forecast is a historical experiment from December 2024, not a current 2026 operational forecast.')

page_before(14)
body(15,'Seasonal naive forecasting repeats the last observed seasonal cycle and provides an essential baseline. ARIMA models serial dependence after differencing; additive Holt-Winters models changing level, trend and annual seasonality. Weather-informed regression adds explanatory variables but requires future weather inputs. Forecasting Principles and Practice [1] explains these methods and chronological evaluation; statsmodels [2] provides the implemented estimators and diagnostics. Austin Water publishes monthly billing aggregates [3, 4], while Open-Meteo supplies historical ERA5 reanalysis [5]. The gap addressed here is the connection between reproducible forecasting, explainable alert screening and explicit allocation feasibility. A forecast alone does not determine how scarce water should be distributed.')
body(17,'The public source snapshots contain 26,819 residential rows and 35,065 commercial rows from January 2012 through October 2025. The analysis uses the complete 2012-2024 calendar years. Water metadata identifies both datasets as Public Domain. Open-Meteo weather is CC BY 4.0; daily mean temperature and rainfall are transformed into monthly mean temperature and rainfall totals. The original snapshots, metadata, byte sizes and SHA-256 hashes are retained for reproducibility.')
anchor=table_after(original[17],['Input or field','Meaning and transformation'],[
('year_month','Billing month YYYYMM, parsed to month start'),('postal_code and customer_class','Record key and retail sector mapping'),('total_gallons','Monthly US gallons; aggregated and divided by 1,000,000'),('Temperature and rainfall','Daily ERA5 data at Austin; monthly mean C and total mm'),('Calendar attributes','Month, year, days in month and weekend days'),('Analysis table','156 monthly rows; four sectors and aggregate demand')],[2.0,4.9])
body(18,'Functional requirements are deterministic validation, exploration, ADF and ACF/PACF diagnostics, at least two forecasting methods, MAE/RMSE/MAPE evaluation, anomaly flags, capacity-constrained allocation and interactive CSV exports. Non-functional requirements are reproducibility, readable decisions, explicit units and no hidden assumptions. Python 3.12, the pinned requirements and a modern browser are required. A 64-bit laptop with about 4 GB free memory is recommended; no dedicated GPU, database or paid API is required. Internet is needed only to install dependencies or refresh public sources.')

page_before(19)
body(20,'The system has six modules: source snapshots, monthly preprocessing, exploratory diagnostics, chronological model comparison, causal anomaly screening and resource allocation. The dashboard reads saved results rather than refitting models on every interaction. This separates the analytical experiment from the scenario controls and allows an offline demonstration after setup.')
anchor=figure(original[20],Path('results/figures/architecture.png'),'Figure 1  Implemented system architecture')
body(21,'ETL extracts the source snapshots, transforms dates and billing classes, aggregates gallons by month and sector, joins monthly weather and loads a processed CSV. Exact duplicates are removed. Conflicting keys, invalid dates and nonfinite volumes stop the pipeline. Two negative commercial source rows are retained as billing adjustments; monthly sector totals are checked for nonnegative values. No duplicate rows or missing sector months occur in this snapshot. Valid unusual demand is retained. If another snapshot contains a missing month, causal forward filling is allowed only after the series starts; leading gaps require repair. This project uses flat files and does not claim to implement a data warehouse.')
body(22,'Training covers 2012-2022, validation covers 2023 and the final test is 2024. Training-only ADF and ACF/PACF support differencing and candidate identification. Converged ARIMA p,q values from 0 to 2 are compared by training AIC. Four model families are ranked by validation RMSE, and the chosen family remains fixed for test evaluation. The weather regression uses history-derived monthly climatology for future weather inputs, avoiding use of actual holdout weather. Final models refit through December 2024 and forecast all 12 months of 2025.')

page_before(23)
body(24,'Algorithm steps are shown below. The main data structure is a date-indexed pandas DataFrame. Preprocessing is approximately O(N log N) for grouping/sorting N source records and O(N) memory. ARIMA fitting cost depends on the optimizer, order and series length; the search fits nine small-order candidates. Forecasting and anomaly screening are linear in the number of monthly observations for fixed order/window. Allocation solves a four-variable linear program per scenario; no general constant-time guarantee is claimed for linear programming.')
anchor=original[24]
steps=[
'1  Validate source columns, dates and record keys; remove exact duplicate rows.',
'2  Map retail classes; exclude utility and wholesale; aggregate monthly sector gallons.',
'3  Join monthly weather and derive calendar features; preserve raw snapshots and hashes.',
'4  Plot sector demand, STL components, rolling mean/variance, ACF and PACF.',
'5  Split chronologically; run training ADF; rank converged ARIMA orders by AIC.',
'6  Forecast the validation year with each model; select the lowest validation RMSE.',
'7  Refit through 2023; predict untouched 2024; compute MAE, RMSE and MAPE.',
'8  Refit through 2024; forecast 2025; split totals using trailing 12-month sector shares.',
'9  Flag |robust seasonal residual score| > 3.5 using prior residuals only.',
'10  Check minimum floors against supply and caps; compare priority and proportional allocations.',
'11  Export metrics, alerts, scenarios and plots; present the interactive dashboard.'
]
for text in steps:anchor=insert_after(anchor,text)
anchor=insert_after(anchor,'Priority allocation objective and constraints',bold=True)
anchor=insert_after(anchor,'Maximize the sum of weighted served demand. For each sector i, minimum_fraction[i] × demand[i] ≤ allocation[i] ≤ min(demand[i], capacity[i]); total allocation ≤ available supply. Minimum fractions are 70%, 35%, 30% and 90%; weights are 4, 2, 1 and 5 for residential, commercial, large volume and municipal respectively. These are declared assumptions. If any floor is infeasible, report emergency service violations rather than a feasible allocation.')

page_before(25)
body(26,'Implementation uses Python, pandas, NumPy, scipy, statsmodels, Matplotlib, Plotly and Streamlit. src/pipeline.py validates and aggregates sources, fits models, evaluates predictions, detects abnormal months and exports evidence. src/allocation.py implements the four-sector service constraints and both allocation policies. app.py supplies six dashboard tabs and adjustable month, supply, capacity and policy controls. tests/ contains numerical and dashboard integration checks. The project uses CSV/JSON files; no database credentials or API keys are needed.')
body(27,'Run the following commands from the project folder to install dependencies, reproduce the analysis, run tests and open the dashboard.')
anchor=original[27]
for line in ['python -m venv .venv','.venv\\Scripts\\Activate.ps1','python -m pip install -r requirements.txt','python -m src.pipeline','python -m unittest discover -s tests -v','python -m streamlit run app.py --global.developmentMode false --browser.gatherUsageStats false']:
    anchor=insert_after(anchor,line);anchor.runs[0].font.name='Consolas';anchor.runs[0].font.size=Pt(9)
anchor=insert_after(anchor,'Important implementation decisions',bold=True)
anchor=insert_after(anchor,'ARIMA fitting scales MG to billion gallons for numerical conditioning and checks optimizer convergence, with a deterministic retry when necessary. Predictions are converted back to MG. The anomaly baseline is demand from the same month a year earlier; prior residuals provide a rolling median and MAD. A valid alert is retained in the training data. The dashboard labels ARIMA intervals separately from the selected-model point forecast and labels infeasible scenarios in red.')
anchor=insert_after(anchor,'Code example',bold=True)
anchor=code_after(anchor,'floors = minimum_fraction * demand\nupper = np.minimum(demand, capacity)\nfeasible = supply >= floors.sum() and np.all(upper >= floors)\nresult = linprog(-priority, A_ub=np.ones((1, 4)),\n                 b_ub=[supply], bounds=list(zip(floors, upper)))')

replace(29,'DAV COMPLEX ENGINEERING PROBLEM RESULTS AND VALIDATION')
body(31,'All 10 automated tests passed. Testing covered supply and capacity limits, minimum service, invalid inputs, duplicate records, zero demand, forecast metrics, future-weather leakage and dashboard controls. Predictions were evaluated on the 2024 test set, which was excluded from model selection.')
t=source_tables[0]
testrows=[('Supply 0-150, both policies','Allocated total ≤ supply; caps respected','All tested allocations satisfy bounds','Pass'),('Supply below service floors','Infeasible emergency with violations','Explicit emergency status and violations','Pass'),('Capacity below a floor','Infeasible even with ample supply','Capacity infeasibility detected','Pass'),('Zero sector demand','Zero allocation without division failure','Zero allocation returned','Pass'),('Negative or NaN supply','Reject invalid input','ValueError raised','Pass'),('Exact and conflicting duplicates','Remove exact; reject conflicting key','Both behaviors verified','Pass'),('Invalid date','Reject malformed source record','ValueError raised','Pass'),('Zero actual in MAPE','Exclude zeros; all-zero MAPE undefined','20% for nonzero example; None for all zero','Pass'),('Dashboard controls','Six tabs; feasible at 75%; emergency at 35%','No application exceptions; statuses verified','Pass')]
for i,row in enumerate(testrows,1):
    cells=t.rows[i].cells if i<len(t.rows) else t.add_row().cells
    for c,text in zip(cells,row):c.text=text
style_table(t,[1.7,2.0,2.4,.65])

page_before(32)
body(33,'Holt-Winters has the lowest validation RMSE and therefore remains the selected planning model. The 2024 test shows that model rankings vary: ARIMA has the lowest test RMSE, while weather/calendar regression has the lowest test MAPE. Selecting a different winner after inspecting the test would weaken the holdout evidence. All reported errors are calculated from actual observations; MG denotes million US gallons per month.')
rows=[(str(r.model),f'{r.validation_RMSE:.2f}',f'{r.test_MAE:.2f}',f'{r.test_RMSE:.2f}',f'{r.test_MAPE:.2f}') for _,r in metrics.iterrows()]
anchor=table_after(original[33],['Model','Validation RMSE','Test MAE','Test RMSE','Test MAPE %'],rows,[2.25,1.25,1.0,1.0,1.2])
anchor=figure(anchor,Path('results/figures/test_forecast.png'),'Figure 2  Forecasts compared with untouched 2024 demand')
body(34,'Training ADF gives p=0.9434 on levels and p=1.12 × 10^-16 after first differencing, supporting d=1 under the test assumptions. The training AIC search selects ARIMA (2,1,2). Residual Ljung-Box p=0.0341 at lag 12 suggests remaining dependence and motivates seasonal-model extensions. Temperature and demand have correlation 0.748 over the analysis period; this descriptive association is not a causal effect.')
body(35,'The experiment covers thirteen years and evaluates one full held-out year. It does not establish universal model superiority. Billing schedules and corrections may differ from physical usage timing. Population is unavailable, future weather is approximated by historical climatology, industrial end-use is unverified, and no actual network or supply data is available. Point forecasts use stable trailing sector shares; uncertainty in those shares is not quantified.')

anchor=insert_after(original[35],'Seasonality and stationarity evidence',bold=True,pagebreak=True)
anchor=figure(anchor,Path('results/figures/consumption.png'),'Figure 3  Selected retail consumption by measured billing class')
anchor=figure(anchor,Path('results/figures/acf_pacf.png'),'Figure 4  Training ACF and PACF after first differencing')
anchor=insert_after(anchor,'The seasonal demand cycle is visible in the sector chart. ACF and PACF show multiple dependencies, so the order is confirmed by the converged AIC grid rather than inferred from a single spike. Rolling mean and variance are provided in the dashboard and results/figures/stationarity.png.')

anchor=insert_after(anchor,'Abnormal consumption and possible explanations',bold=True,pagebreak=True)
anchor=figure(anchor,Path('results/figures/anomalies.png'),'Figure 5  Four months flagged for review by a causal seasonal residual screen')
flagged=anomalies.loc[anomalies.flag]
rows=[(r.date[:7],f'{r.actual_mg:.2f}',f'{r.robust_score:.2f}',f'{r.temperature_c:.1f}',f'{r.rainfall_mm:.1f}') for _,r in flagged.iterrows()]
anchor=table_after(anchor,['Month','Actual MG','Robust score','Mean C','Rain mm'],rows,[1.3,1.3,1.3,1.3,1.3])
anchor=insert_after(anchor,'Each month is compared with its year-earlier demand. The detector uses only prior residuals to estimate a 24-month median and MAD, with at least 12 residuals required. Initial months are unscored. A score beyond ±3.5 triggers investigation. Review billing corrections, drought restrictions, leaks, unusual events and weather. Available weather can suggest associations, but no flagged month is a confirmed leak. Without incident labels, anomaly precision and recall cannot be reported.')
anchor=insert_after(anchor,'Year-earlier weather comparisons provide limited explanations. October 2019 demand rose by 1,268.88 MG while rainfall fell by 69.8 mm, consistent with increased irrigation as a hypothesis. August 2024 demand fell by 1,052.64 MG while mean temperature fell by 2.24 C, potentially reducing outdoor demand. October 2015 was substantially wetter despite higher demand, and August 2022 had only modest weather changes; those increases need billing, restriction or incident records to explain.')

anchor=insert_after(anchor,'Alternative distribution scenarios',bold=True,pagebreak=True)
anchor=insert_after(anchor,'For January 2025 the selected total forecast is 3,394.65 MG. At 75% supply, available water is 2,545.99 MG and the supply-demand gap is 848.66 MG. The following comparison uses 100% sector capacity, matching the dashboard default. Both policies meet the assumed minimum floors; their remaining allocations differ.')
from src.allocation import allocate,SECTORS
demand=forecast.loc[0,SECTORS].to_numpy(dtype=float)
priority,_=allocate(demand,float(forecast.loc[0,'selected_forecast_mg']*.75),'Priority')
proportional,_=allocate(demand,float(forecast.loc[0,'selected_forecast_mg']*.75),'Proportional')
rows=[(s,f'{demand[i]:.2f}',f'{priority.minimum_mg.iloc[i]:.2f}',f'{priority.allocated_mg.iloc[i]:.2f}',f'{proportional.allocated_mg.iloc[i]:.2f}') for i,s in enumerate(SECTORS)]
anchor=table_after(anchor,['Sector','Demand MG','Floor MG','Priority MG','Proportional MG'],rows,[1.6,1.2,1.2,1.35,1.35])
anchor=insert_after(anchor,'At 35% supply, only 1,188.13 MG is available while the assumed minimum floors require approximately 1,956.88 MG. The system therefore reports infeasibility and an emergency fallback, including floor violations. Additional water, temporary distribution, demand restrictions or restored capacity are required. A sector cap below its floor also makes service infeasible even when aggregate supply is sufficient.')

anchor=insert_after(anchor,'Decision support demonstration',bold=True,pagebreak=True)
anchor=figure(anchor,Path('docs/dashboard_allocation.png'),'Figure 6  Dashboard showing the 75 percent supply scenario',width=6.8)
anchor=insert_after(anchor,'The dashboard presents forecasts, model errors, weather patterns, stationarity tests, unusual demand and allocation results. Users can change the planning month, supply, sector capacity and allocation policy, then download the results as CSV files.')

anchor=insert_after(anchor,'Notebook code and execution outputs',bold=True,pagebreak=True)
anchor=insert_after(anchor,'notebooks/Urban_Water_Forecasting_Colab.ipynb contains the executable analysis and saved outputs. In Google Colab, select Runtime then Run all. Setup clones the public project and installs its dependencies. Subsequent cells reproduce preprocessing, forecasts, ADF diagnostics, anomaly alerts and allocation scenarios from the source snapshots. The submitted outputs were executed locally with a Jupyter kernel; the notebook includes setup for Colab execution.')
anchor=code_after(anchor,'from src.pipeline import run\nrun()\ncomparison = pd.read_csv("results/metrics.csv")\ndisplay(comparison.round(2))\nalerts = pd.read_csv("results/anomalies.csv")\ndisplay(alerts.loc[alerts["flag"]])')
anchor=figure(anchor,Path('docs/notebook_output.png'),'Figure 7  Executed notebook output showing forecast errors and anomaly alerts',width=6.8)
anchor=insert_after(anchor,'The notebook saves tables and color plots in its output cells. results/notebook_outputs.md provides a separate readable record of metrics, ADF results, alerts, allocation comparisons and test output. Re-running uses the same chronological split and validation selection rule as the dashboard.')

page_before(36)
body(37,'')
t=source_tables[1]
while len(t.rows)<14:t.add_row()
for i,row in enumerate(t.rows[1:],1):
    for c,text in zip(row.cells,[f'{i}. ____________________','____________________','____________________________']):c.text=text
style_table(t,[2.4,2.0,2.6])
body(39,'The project implements all eight required analytics tasks using public billing and weather observations. It demonstrates reproducible cleaning, seasonal exploration, stationarity testing, ARIMA and three alternatives, holdout forecasting, abnormal-month screening, two constrained allocation policies and an interactive dashboard. Holt-Winters is selected by validation and achieves 6.70% test MAPE. Allocation explicitly distinguishes feasible service from emergency violations. Future work should add verified daily or zone-level metering, actual supply and hydraulic capacities, critical-facility service requirements, population data, seasonal ARIMA and rolling-origin evaluation. Real deployment also requires calibrated uncertainty and validation with municipal operators.')
anchor=insert_after(original[39],'References',bold=True,pagebreak=True)
refs=[
'[1] Hyndman RJ and Athanasopoulos G. Forecasting Principles and Practice, 3rd edition. https://otexts.com/fpp3/  Methods, seasonal baselines and time-series evaluation.',
'[2] statsmodels documentation. ARIMA, ExponentialSmoothing, adfuller, ACF/PACF and Ljung-Box diagnostics. https://www.statsmodels.org/stable/api.html',
'[3] City of Austin. Austin Water Residential Water Consumption. Dataset sxk7-7k6z. https://data.austintexas.gov/d/sxk7-7k6z  Public Domain according to source metadata.',
'[4] City of Austin. Austin Water Commercial Water Consumption. Dataset 5h9c-wmds. https://data.austintexas.gov/d/5h9c-wmds  Public Domain according to source metadata.',
'[5] Open-Meteo Historical Weather API. ERA5 reanalysis for Austin. https://open-meteo.com/en/docs/historical-weather-api  Weather data CC BY 4.0; daily data transformed to monthly aggregates.',
'[6] SciPy documentation. scipy.optimize.linprog. https://docs.scipy.org/doc/scipy/reference/generated/scipy.optimize.linprog.html  Linear optimization with bound and supply constraints.',
]
for ref in refs:anchor=insert_after(anchor,ref)

# Preserve source geometry, headers, footers, images, opaque parts and style definitions.
# Python-docx is used to add content; reference preserve-only parts are restored byte-for-byte.
final=ROOT/'docs/CEP_Report.docx'
d.save(final)
with ZipFile(template) as src,ZipFile(final) as generated:
    parts={n:generated.read(n) for n in generated.namelist()}
    editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml'}
    for n in src.namelist():
        if n not in editable:parts[n]=src.read(n)

# Black report typography and tables; retain color in pictures and code.
from lxml import etree as ET
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
V='urn:schemas-microsoft-com:vml'
monochrome_parts=set()
for name,content in list(parts.items()):
    if not name.startswith('word/') or not name.endswith('.xml'):continue
    xml=ET.fromstring(content)
    for color in xml.iter('{'+W+'}color'):
        run=color.getparent().getparent()
        if run.tag=='{'+W+'}r' and run.xpath('.//*[local-name()="rFonts"][@*[local-name()="ascii"]="Consolas"]'):continue
        for attr in list(color.attrib):
            if ET.QName(attr).localname.startswith('theme'):del color.attrib[attr]
        color.set('{'+W+'}val','000000')
    for shade in xml.iter('{'+W+'}shd'):
        for attr in list(shade.attrib):
            if ET.QName(attr).localname.startswith('theme'):del shade.attrib[attr]
        shade.set('{'+W+'}fill','FFFFFF');shade.set('{'+W+'}color','000000')
    for border in xml.xpath('//*[local-name()="tblBorders" or local-name()="tcBorders"]/*'):
        border.set('{'+W+'}color','000000')
    for blip in xml.iter('{'+A+'}blip'):
        for effect in list(blip.findall('{'+A+'}grayscl')):blip.remove(effect)
    for image in xml.iter('{'+V+'}imagedata'):image.attrib.pop('grayscale',None)
    new=ET.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    if new!=content:parts[name]=new;monochrome_parts.add(name)
# Preserve full-resolution figure media in Word rather than automatic downsampling.
settings=ET.fromstring(parts['word/settings.xml'])
if settings.find('{'+W+'}doNotAutoCompressPictures') is None:
    ET.SubElement(settings,'{'+W+'}doNotAutoCompressPictures')
ns14='http://schemas.microsoft.com/office/word/2010/wordml'
dpi=settings.find('{'+ns14+'}defaultImageDpi')
if dpi is None:dpi=ET.SubElement(settings,'{'+ns14+'}defaultImageDpi')
dpi.set('{'+ns14+'}val','600')
parts['word/settings.xml']=ET.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
monochrome_parts.add('word/settings.xml')
with ZipFile(final,'w',ZIP_DEFLATED) as z:
    for name,content in parts.items():z.writestr(name,content)
with ZipFile(template) as src,ZipFile(final) as finalzip:
    mismatches=[n for n in src.namelist() if n not in editable|monochrome_parts and src.read(n)!=finalzip.read(n)]
assert not mismatches,mismatches
assert len(Document(final).sections)==len(Document(template).sections)
print('Report saved:',final)
print('Reference SHA-256:',hashlib.sha256(template.read_bytes()).hexdigest())
print('Template geometry preserved; black typography and tables, color visuals and code.')
